"""
Top-level orchestrator for funeral / memorial bulletin generation.

Parallels ``BulletinBuilder`` (the Sunday/special-service builder) but
on the YAML-driven side: it consumes a per-service ``FuneralData``
dataclass + the existing scripture cache + the existing song-lookup
pipeline, and produces a Document.

Usage
=====

    from bulletin.document.funeral_builder import FuneralBuilder
    from bulletin.sources.funeral_data import load_service

    fd = load_service("2026-01-31-cox")
    builder = FuneralBuilder(fd, scripture_readings, song_lookup_fn)
    doc = builder.build()
    doc.save("output/2026-01-31 - Burial of the Dead - Annette Cox.docx")

The CLI driver (``generate.py --funeral <slug>``) wraps this with a
small helper that fetches the readings and supplies the song lookup.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from docx.oxml import parse_xml

from bulletin.document.styles import configure_document
from bulletin.document.templates import (
    _replace_all_placeholders,
    _pin_floating_shapes_to_first_paragraph,
    append_back_cover,
    setup_footers,
)
from bulletin.document.sections.burial import (
    add_burial_service, _format_clergy_title,
)
from bulletin.sources.funeral_data import FuneralData


_TEMPLATES_DIR = Path(__file__).resolve().parent.parent.parent / "templates"
_FUNERAL_COVER_TEMPLATE = "front_cover_funeral.docx"
_FUNERAL_BACK_COVER_TEMPLATE = "back_cover_funeral.docx"

_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


# ---------------------------------------------------------------------------
# Participants — formatted paragraphs that replace the {{PARTICIPANTS}}
# placeholder inside the back cover's text box.
# ---------------------------------------------------------------------------

# Role rows in the order they should appear on the back cover.
# Each tuple: (yaml-key, display-label, is_list, expects_clergy_title)
_PARTICIPANT_ROWS = [
    ("celebrant",       "Celebrant",         False, True),
    ("assisting",       "Assisting Priests", True,  True),
    ("deacon",          "Deacon",            False, True),
    ("readers",         "Readers",           True,  False),
    ("crucifer",        "Crucifer",          False, False),
    ("chalice_bearers", "Chalice Bearers",   True,  False),
    ("musician",        "Musician",          False, False),
    ("vocalist",        "Vocalist",          False, False),
    ("pall_bearers",    "Pall Bearers",      True,  False),
]


def _participant_rows(participants: dict) -> list[tuple[str, str]]:
    """Resolve the participants dict into an ordered list of (name(s),
    label) pairs, dropping empty entries.

    For multi-person roles (Assisting Priests, Readers, Chalice Bearers,
    Pall Bearers) the names are joined on a single line separated by
    commas — the existing convention for the inline participants page.
    Clergy roles get "The Rev." prepended via _format_clergy_title when
    the YAML doesn't already include a title.
    """
    rows: list[tuple[str, str]] = []
    for key, label, is_list, clergy in _PARTICIPANT_ROWS:
        val = participants.get(key)
        if not val:
            continue
        if is_list:
            names = [n for n in val if n]
            if not names:
                continue
            if clergy:
                names = [_format_clergy_title(n) for n in names]
            rows.append((", ".join(names), label))
        else:
            name = _format_clergy_title(val) if clergy else val
            rows.append((name, label))
    return rows


def _participant_paragraph_xml(text: str, style_id: str) -> str:
    """Build a single <w:p> using the template's Staff-Name / Staff-
    Title styles. These already carry bold/italic/centering — the
    runs only need to carry the text.
    """
    # The text is XML-escaped because it may contain ampersands or
    # angle brackets in names like "Burns & Sons" or "<unknown>".
    from xml.sax.saxutils import escape
    safe = escape(text)
    return (
        f'<w:p xmlns:w="{_W_NS}">'
        f'<w:pPr><w:pStyle w:val="{style_id}"/></w:pPr>'
        f'<w:r><w:t xml:space="preserve">{safe}</w:t></w:r>'
        f'</w:p>'
    )


def _inject_participants(doc: "Document", participants: dict) -> None:
    """Replace every ``{{PARTICIPANTS}}`` placeholder paragraph in the
    document with a sequence of formatted participant paragraphs.

    Each role becomes two paragraphs:
        Staff-Name   — the name(s), bold + centered.
        Staff-Title  — the role label, italic + centered.

    The Staff-Name style already supplies the vertical breathing
    room via its ``space_before``, so no blank spacer paragraphs
    are inserted between role blocks.

    The placeholder paragraph's parent is unchanged — the new
    paragraphs simply take its slot. Word can split a placeholder
    across runs after a spell-check pass, so the matcher
    concatenates all the paragraph's run text before testing.
    """
    from docx.oxml.ns import qn

    rows = _participant_rows(participants)
    if not rows:
        return

    # Build the replacement paragraphs — Staff-Name carries
    # space-before, so adjacent blocks already get visual breathing
    # room without empty intervening paragraphs.
    new_paragraph_xml: list[str] = []
    for name, label in rows:
        new_paragraph_xml.append(_participant_paragraph_xml(name, "Staff-Name"))
        new_paragraph_xml.append(_participant_paragraph_xml(label, "Staff-Title"))

    # Locate every paragraph that contains the placeholder text.
    # iter() walks the entire tree, including any text-box content
    # if the template uses one.
    target_paragraphs = []
    for para in doc.element.iter(qn("w:p")):
        full_text = "".join(
            (t.text or "") for t in para.iter(qn("w:t"))
        )
        if "{{PARTICIPANTS}}" in full_text:
            target_paragraphs.append(para)

    for placeholder_p in target_paragraphs:
        parent = placeholder_p.getparent()
        index = list(parent).index(placeholder_p)
        for offset, xml in enumerate(new_paragraph_xml):
            parent.insert(index + offset, parse_xml(xml))
        parent.remove(placeholder_p)


def _pin_back_cover_logo(doc: "Document") -> None:
    """Pin the back-cover floating shape(s) to the first paragraph of
    the back cover section.

    The funeral back cover places a parish logo as an absolutely-
    positioned floating image near the bottom of the page
    (``positionV relativeFrom="page"``). Word's rule is that
    a floating shape renders on whichever physical page its anchor
    paragraph sits on. Originally the logo is anchored to a
    paragraph somewhere inside the back-cover content; as the
    participants list grows, that anchor paragraph drifts down the
    page, and if the participants ever push the anchor onto a
    second page the logo follows.

    Moving the anchor's parent ``<w:r>`` to the FIRST paragraph of
    the back-cover section guarantees the logo always renders on
    the back cover, since that paragraph is the start of a
    "next-page" section and is always on a fresh page. The visual
    Y position of the logo doesn't change — that's still locked to
    the page via ``relativeFrom="page"``.
    """
    from docx.oxml.ns import qn

    body = doc.element.body
    # Find the last in-paragraph <w:sectPr> — that's the section
    # break that closes the bulletin content and starts the back
    # cover. Skip any sectPr in the body itself (the trailing one),
    # since that defines the LAST section's page setup, not a break.
    last_inline_sect_pr_para = None
    for p in body.iterchildren(qn("w:p")):
        for pPr in p.iterchildren(qn("w:pPr")):
            if pPr.find(qn("w:sectPr")) is not None:
                last_inline_sect_pr_para = p

    if last_inline_sect_pr_para is None:
        # No section break — fallback: pin to body's first paragraph.
        first_back_para = next(body.iterchildren(qn("w:p")), None)
    else:
        # The first paragraph AFTER the section break is the
        # back cover's first paragraph.
        first_back_para = None
        seen_break = False
        for p in body.iterchildren(qn("w:p")):
            if seen_break:
                first_back_para = p
                break
            if p is last_inline_sect_pr_para:
                seen_break = True

    if first_back_para is None:
        return

    # Collect every floating-shape run that lives AFTER the back
    # cover's first paragraph (i.e., in the back-cover section).
    runs_to_move = []
    in_back_cover = False
    for p in body.iterchildren(qn("w:p")):
        if p is first_back_para:
            in_back_cover = True
            continue
        if not in_back_cover:
            continue
        for run in p.iter(qn("w:r")):
            for drawing in run.iter(qn("w:drawing")):
                if drawing.find(qn("wp:anchor")) is not None:
                    runs_to_move.append(run)
                    break

    for run in runs_to_move:
        run.getparent().remove(run)
        first_back_para.append(run)


class FuneralBuilder:
    """Builds a funeral / memorial bulletin .docx from a ``FuneralData``.

    Args:
        fd:                  A loaded ``FuneralData``.
        scripture_readings:  Dict ``{reference: ScriptureReading}`` for
                             every reading the YAML names — the CLI
                             driver fetches these via the existing
                             scripture cache before constructing the
                             builder.
        song_lookup_fn:      ``callable(title, service) -> dict | None``
                             — same shape as the Sunday pipeline.
        report:              Optional ``RunReport`` to record warnings.
    """

    def __init__(self, fd: FuneralData, scripture_readings: dict,
                 song_lookup_fn, report=None):
        self.fd = fd
        self.scripture = scripture_readings
        self.song_lookup = song_lookup_fn
        self.report = report

    # ------------------------------------------------------------------
    # build()
    # ------------------------------------------------------------------

    def build(self) -> Document:
        """Open the cover template, substitute placeholders, render the
        liturgy, append the back cover, and return the Document."""
        doc = self._load_cover()
        configure_document(doc)

        add_burial_service(doc, self.fd, self.scripture, self.song_lookup)

        # Append the funeral-specific back cover (BCP boilerplate +
        # parish logo + a text-boxed Participants block). The text
        # box's {{PARTICIPANTS}} placeholder is substituted with
        # formatted paragraphs (bold Name / italic role, alternating)
        # rather than a single text value, so we run the participant
        # injector after append_back_cover has copied the template
        # content into the main doc.
        append_back_cover(doc, template_name=_FUNERAL_BACK_COVER_TEMPLATE)
        _inject_participants(doc, self.fd.participants)
        _pin_back_cover_logo(doc)

        setup_footers(
            doc,
            date_str=self.fd.service_date_long,
            service_time="",  # funerals don't print a service-time footer
            liturgical_title=self._footer_title(),
        )
        return doc

    # ------------------------------------------------------------------
    # Cover template handling
    # ------------------------------------------------------------------

    def _load_cover(self) -> Document:
        """Load the two-page funeral cover template and substitute the
        five text placeholders (``{{DATE}}``, ``{{SUBTITLE}}``,
        ``{{NAME}}``, ``{{LIFE_DATES}}``, ``{{BIO}}``).

        The photo placeholder is an embedded image in the template that
        the user swaps out manually in the post-generation edit pass.
        Services with no bio AND no photo: the user deletes the inner
        page in the same edit pass — simpler than conditional template
        assembly.
        """
        path = _TEMPLATES_DIR / _FUNERAL_COVER_TEMPLATE
        if not path.exists():
            raise FileNotFoundError(
                f"Funeral cover template not found: {path}\n"
                f"Expected at: {path}"
            )
        doc = Document(str(path))

        bio = self.fd.deceased.get("bio") or ""
        # Multi-paragraph bios: collapse to a single string with double
        # spaces between paragraphs. The user's hand-edit pass splits
        # them back into paragraphs in Word — supporting true multi-
        # paragraph substitution requires custom XML manipulation that
        # we'll add in a follow-up if it turns out to be tedious.
        bio_collapsed = "  ".join(p.strip() for p in bio.split("\n\n") if p.strip())

        replacements = {
            "{{DATE}}":       self.fd.service_date_long,
            "{{SUBTITLE}}":   self.fd.cover_subtitle_resolved,
            "{{NAME}}":       self.fd.deceased.get("full_name", ""),
            "{{LIFE_DATES}}": self.fd.life_dates,
            "{{BIO}}":        bio_collapsed,
        }
        _replace_all_placeholders(doc, replacements)
        _pin_floating_shapes_to_first_paragraph(doc)
        return doc

    def _footer_title(self) -> str:
        """Footer left text — matches the convention in the existing
        bulletins: 'Burial of the Dead, Rite II' / 'Memorial Service'
        / etc."""
        kind = self.fd.service_kind
        rite = f", Rite {self.fd.rite}"
        if kind == "burial":
            return f"Burial of the Dead{rite}"
        if kind == "memorial":
            return f"Memorial Service{rite}"
        if kind == "committal":
            return f"Graveside Service{rite}"
        return f"Funeral Service{rite}"

    # ------------------------------------------------------------------
    # Output filename
    # ------------------------------------------------------------------

    def output_filename(self) -> str:
        """Suggested filename for the generated .docx — mirrors the
        convention used by the existing Pages bulletins, e.g.
        '2026-01-31 - Burial of the Dead - Annette Cox.docx'."""
        kind_label = {
            "burial":   "Burial of the Dead",
            "memorial": "Memorial Service",
            "committal": "Graveside Service",
        }.get(self.fd.service_kind, "Funeral Service")
        date = self.fd.service["date"]
        if hasattr(date, "isoformat"):
            date = date.isoformat()
        first = self.fd.deceased.get("preferred_name", "") or ""
        last = self.fd.deceased.get("full_name", "").split()[-1]
        name = f"{first} {last}".strip() or self.fd.deceased.get("full_name", "")
        return f"{date} - {kind_label} - {name}.docx"
